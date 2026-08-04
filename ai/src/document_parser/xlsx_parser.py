"""XLSX → Markdown. 첫 시트만 읽는다 (v1 범위 — 다중 시트는 근거 없음)."""

import datetime
from pathlib import Path

import openpyxl

from .models import PageResult, ParseError, ParseResult
from .tabular_markdown import rows_to_markdown


def parse_xlsx(path: Path) -> ParseResult:
    try:
        workbook = openpyxl.load_workbook(path, data_only=True)
    except Exception:
        return ParseResult(
            error=ParseError("corrupt_document", "XLSX 파일을 읽을 수 없습니다.")
        )

    sheet = workbook[workbook.sheetnames[0]]
    warnings = tuple(
        f"시트 '{name}'는 무시했습니다 — 첫 시트만 지원합니다."
        for name in workbook.sheetnames[1:]
    )

    rows = _sheet_rows(sheet)

    try:
        markdown = rows_to_markdown(rows)
    except ValueError:
        return ParseResult(
            error=ParseError("not_tabular", "표로 인식할 수 없는 XLSX입니다.")
        )

    return ParseResult(
        text=markdown,
        pages=(PageResult(page=1, text=markdown, method="native", quality_score=1.0),),
        quality_score=1.0,
        warnings=warnings,
    )


def _sheet_rows(sheet) -> list[list[str]]:
    merged_lookup = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_lookup[(row, col)] = top_left

    rows = []
    for row in sheet.iter_rows():
        cells = [_cell_text(cell, merged_lookup) for cell in row]
        rows.append(cells)
    return rows


def _cell_text(cell, merged_lookup: dict) -> str:
    value = merged_lookup.get((cell.row, cell.column), cell.value)
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return _format_datetime(value, cell.number_format or "")
    return str(value)


def _format_datetime(value, number_format: str) -> str:
    """값의 타입이 아니라 셀 서식으로 날짜/시간 표시 여부를 정한다.

    `datetime.time(0, 0)`은 파이썬에서 falsy 가 아니라서, 값 타입만 보면 날짜만
    있는 셀(내부적으로 자정 시각이 채워진 `datetime.datetime`)에도 "00:00"이
    잘못 붙는다 — 스파이크에서 실측한 버그다. 서식 문자열에 실제로 시간 토큰(`h`)이
    있는지로 판단해야 정확하다.
    """
    fmt = number_format.lower()
    has_time = "h" in fmt
    has_date = any(token in fmt for token in "ymd")
    if has_date and has_time:
        return value.strftime("%Y-%m-%d %H:%M")
    if has_time:
        return value.strftime("%H:%M")
    return value.strftime("%Y-%m-%d")
