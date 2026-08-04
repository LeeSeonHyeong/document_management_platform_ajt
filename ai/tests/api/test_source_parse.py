"""원본문서 파싱 API — `POST /internal/v1/source-parses` (계약 v1.3.0).

Spring Boot 가 현재 호출하는 유일한 AI 엔드포인트다. 파싱 자체는
`document_parser` 가 이미 하므로, 여기서 보는 것은 **계약 경계**다: multipart 를
받아 파서에 넘기고, 파서의 결과·오류를 계약이 정한 응답과 오류 코드로 바꾼다.
"""

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from wiki_api.app import create_app

API_KEY = "secret-key"
PATH = "/internal/v1/source-parses"

MARKDOWN = "# 취업 규칙\n\n## 3장 휴가\n\n연차는 15일을 부여한다.\n"


def _client():
    return TestClient(create_app(api_key=API_KEY), raise_server_exceptions=False)


def _form(**overrides):
    form = {
        "requestId": "parse-request-1",
        "sourceType": "wiki",
        "sourceId": "15",
        "originalFileName": "취업규칙.md",
        "mimeType": "text/markdown",
    }
    form.update(overrides)
    return form


def _post(content: bytes = MARKDOWN.encode(), filename: str = "취업규칙.md", **overrides):
    return _client().post(
        PATH,
        data=_form(**overrides),
        files={"file": (filename, io.BytesIO(content), "application/octet-stream")},
        headers={"X-Internal-API-Key": API_KEY},
    )


# ----- 성공 경로 ---------------------------------------------------------------


def test_a_markdown_source_comes_back_as_markdown():
    response = _post()

    assert response.status_code == 200, response.text
    body = response.json()
    assert body["requestId"] == "parse-request-1"
    assert body["sourceType"] == "wiki"
    assert body["sourceId"] == "15"
    assert "연차는 15일을 부여한다" in body["parsedMarkdown"]
    assert body["warnings"] == []


def test_the_response_carries_only_the_contract_fields():
    """Spring 의 `SourceParseResponse` 가 다섯 필드를 전부 비어 있지 않게 검증한다 —
    필드가 늘거나 이름이 달라지면 그쪽에서 `INVALID_RESPONSE` 가 된다."""
    body = _post().json()

    assert set(body) == {"requestId", "sourceType", "sourceId", "parsedMarkdown",
                         "warnings"}


def test_a_docx_source_keeps_its_heading_structure(tmp_path):
    """DOCX 헤딩이 Markdown 헤더로 살아 있어야 변환 에이전트가 절 구조를 본다."""
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("휴가 규정", level=1)
    document.add_paragraph("연차는 15일이다.")
    document.save(path)

    response = _post(path.read_bytes(), "policy.docx",
                     originalFileName="policy.docx",
                     mimeType="application/vnd.openxmlformats-officedocument"
                              ".wordprocessingml.document")

    assert response.status_code == 200, response.text
    assert response.json()["parsedMarkdown"].startswith("# 휴가 규정")


def test_the_source_type_is_echoed_back():
    body = _post(sourceType="schedule", sourceId="group-7").json()

    assert body["sourceType"] == "schedule"
    assert body["sourceId"] == "group-7"


def test_a_schedule_csv_source_is_accepted():
    csv_bytes = ("제목,시작,종료\n"
                 "전사 워크샵,2026-08-05 14:00,2026-08-05 18:00\n").encode("utf-8")

    response = _post(csv_bytes, "일정.csv", sourceType="schedule", sourceId="group-7",
                     originalFileName="일정.csv", mimeType="text/csv")

    assert response.status_code == 200, response.text
    assert "전사 워크샵" in response.json()["parsedMarkdown"]


def test_a_schedule_xlsx_source_is_accepted(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.append(["제목"])
    wb.active.append(["전사 워크샵"])
    path = tmp_path / "일정.xlsx"
    wb.save(path)

    response = _post(path.read_bytes(), "일정.xlsx", sourceType="schedule",
                     sourceId="group-7", originalFileName="일정.xlsx",
                     mimeType="application/vnd.openxmlformats-officedocument"
                              ".spreadsheetml.sheet")

    assert response.status_code == 200, response.text
    assert "전사 워크샵" in response.json()["parsedMarkdown"]


def test_a_wiki_csv_source_is_still_rejected():
    """FR-DOC-002 — 위키 원본은 CSV·XLSX 를 허용하지 않는다. 일정만 열었다."""
    response = _post(b"a,b\n1,2\n", "문서.csv", sourceType="wiki",
                     originalFileName="문서.csv", mimeType="text/csv")

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_SOURCE_PARSE_REQUEST"


def test_a_not_tabular_csv_is_a_bad_request_not_a_parse_failure():
    """공백만 있는 CSV(빈 파일은 아니라 별도 검증에 걸리지 않는다)는 표로 인식될
    내용이 없다 — 요청 자체의 문제이므로 400 이어야 한다 (Task 7 에서
    `encoding_undetected`용으로 추가한 테스트와 짝을 이룬다)."""
    response = _post(b"   \n  \n", "일정.csv", sourceType="schedule", sourceId="group-7",
                     originalFileName="일정.csv", mimeType="text/csv")

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_SOURCE_PARSE_REQUEST"


def test_an_encoding_failure_is_a_bad_request_not_a_parse_failure():
    """인코딩을 못 판별한 것은 요청 자체의 문제다 — 500 이 아니라 400 이어야 한다."""
    response = _post(b"\x80\x81\x82\x83", "일정.csv", sourceType="schedule",
                     sourceId="group-7", originalFileName="일정.csv",
                     mimeType="text/csv")

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_SOURCE_PARSE_REQUEST"


# ----- 400: 요청이 잘못된 경우 --------------------------------------------------


def test_an_unsupported_extension_is_a_bad_request():
    response = _post(b"a,b\n1,2\n", "일정.csv", originalFileName="일정.csv",
                     mimeType="text/csv")

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_SOURCE_PARSE_REQUEST"


def test_an_unknown_source_type_is_a_bad_request():
    response = _post(sourceType="invoice")

    assert response.status_code == 400
    body = response.json()
    assert body["code"] == "INVALID_SOURCE_PARSE_REQUEST"
    assert [e["field"] for e in body["fieldErrors"]] == ["sourceType"]


@pytest.mark.parametrize("missing", ["requestId", "sourceId", "originalFileName"])
def test_missing_metadata_is_a_bad_request(missing):
    form = _form()
    del form[missing]
    response = _client().post(
        PATH, data=form,
        files={"file": ("f.md", io.BytesIO(MARKDOWN.encode()), "text/markdown")},
        headers={"X-Internal-API-Key": API_KEY})

    assert response.status_code == 400
    body = response.json()
    assert body["code"] == "INVALID_SOURCE_PARSE_REQUEST"
    assert [e["field"] for e in body["fieldErrors"]] == [missing]


def test_an_empty_file_is_a_bad_request():
    response = _post(b"")

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_SOURCE_PARSE_REQUEST"


def test_the_extension_comes_from_the_declared_name_not_the_upload():
    """Spring 이 검증한 이름은 `originalFileName` 이다. multipart 의 파일명은
    임시 경로일 수 있으므로 그것으로 형식을 판단하면 안 된다."""
    response = _post(MARKDOWN.encode(), "tmp-upload-9182",
                     originalFileName="취업규칙.md")

    assert response.status_code == 200, response.text


# ----- 500: 파싱이 실패한 경우 --------------------------------------------------


def test_a_corrupt_document_is_a_parse_failure():
    """형식은 지원하는데 내용을 읽지 못한 경우다 — 요청 잘못이 아니라 처리 실패라
    계약이 500 `DOCUMENT_PARSE_FAILED` 로 정의한다."""
    response = _post(b"not a real docx", "broken.docx",
                     originalFileName="broken.docx",
                     mimeType="application/vnd.openxmlformats-officedocument"
                              ".wordprocessingml.document")

    assert response.status_code == 500
    assert response.json()["code"] == "DOCUMENT_PARSE_FAILED"


# ----- 401 ---------------------------------------------------------------------


def test_a_wrong_api_key_is_rejected():
    response = _client().post(
        PATH, data=_form(),
        files={"file": ("f.md", io.BytesIO(MARKDOWN.encode()), "text/markdown")},
        headers={"X-Internal-API-Key": "wrong"})

    assert response.status_code == 401
    assert response.json()["code"] == "INVALID_INTERNAL_API_KEY"
