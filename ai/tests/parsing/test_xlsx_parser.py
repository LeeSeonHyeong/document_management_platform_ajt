import datetime
from pathlib import Path

from openpyxl import Workbook

from document_parser.xlsx_parser import parse_xlsx


def _save(tmp_path: Path, build) -> Path:
    wb = Workbook()
    build(wb)
    path = tmp_path / "schedule.xlsx"
    wb.save(path)
    return path


def test_a_single_sheet_table_becomes_a_markdown_table(tmp_path: Path):
    def build(wb):
        ws = wb.active
        ws.title = "일정"
        ws.append(["제목", "장소"])
        ws.append(["전사 워크샵", "본사 대강당"])
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert result.text == (
        "| 제목 | 장소 |\n| --- | --- |\n| 전사 워크샵 | 본사 대강당 |"
    )


def test_a_date_only_cell_has_no_spurious_midnight(tmp_path: Path):
    """실측 버그: 값 타입(datetime.datetime)만 보면 날짜만 있는 셀도 자정이
    붙는다 — `datetime.time(0, 0)`이 파이썬에서 falsy 가 아니기 때문이다."""
    def build(wb):
        ws = wb.active
        ws.title = "일정"
        ws.append(["날짜", "제목"])
        ws.append([datetime.date(2026, 8, 10), "날짜만 있는 셀"])
        ws["A2"].number_format = "yyyy-mm-dd"
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert "2026-08-10 |" in result.text
    assert "00:00" not in result.text


def test_a_datetime_cell_shows_both_date_and_time(tmp_path: Path):
    def build(wb):
        ws = wb.active
        ws.title = "일정"
        ws.append(["시작", "제목"])
        ws.append([datetime.datetime(2026, 8, 5, 14, 0), "워크샵"])
        ws["A2"].number_format = "yyyy-mm-dd hh:mm"
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert "2026-08-05 14:00 |" in result.text


def test_a_merged_cell_repeats_its_value_across_the_range(tmp_path: Path):
    def build(wb):
        ws = wb.active
        ws.title = "안내"
        ws.append(["참가 대상", None])
        ws.merge_cells("A1:B1")
        ws.append(["부서", "인원"])
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert "| 참가 대상 | 참가 대상 |" in result.text


def test_extra_sheets_are_reported_as_warnings_not_parsed(tmp_path: Path):
    def build(wb):
        ws = wb.active
        ws.title = "일정"
        ws.append(["제목"])
        wb.create_sheet("안내")
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert result.warnings == ("시트 '안내'는 무시했습니다 — 첫 시트만 지원합니다.",)


def test_a_row_of_empty_string_cells_is_dropped_like_csv(tmp_path: Path):
    """CSV 의 `,,` 와 동등한 입력 — 셀이 빈 문자열(None 이 아님)인 행도 제외돼야
    한다 (파서 간 동작 통일)."""
    def build(wb):
        ws = wb.active
        ws.title = "일정"
        ws.append(["제목", "시작"])
        ws.append(["", ""])
        ws.append(["워크샵", "2026-08-05"])
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is None
    assert result.text == (
        "| 제목 | 시작 |\n| --- | --- |\n| 워크샵 | 2026-08-05 |"
    )


def test_an_empty_first_sheet_is_not_tabular(tmp_path: Path):
    def build(wb):
        wb.active.title = "일정"
    path = _save(tmp_path, build)

    result = parse_xlsx(path)

    assert result.error is not None
    assert result.error.code == "not_tabular"


def test_a_corrupt_xlsx_is_reported(tmp_path: Path):
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a real xlsx")

    result = parse_xlsx(path)

    assert result.error is not None
    assert result.error.code == "corrupt_document"
