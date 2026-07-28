from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from document_parser import ParseOptions, parse


def _write_policy(path: Path) -> None:
    document = Document()
    document.add_heading("Time off")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Type"
    table.rows[0].cells[1].text = "Minimum"
    table.rows[1].cells[0].text = "Paid leave"
    table.rows[1].cells[1].text = "25 days"
    document.add_paragraph("Manager approval is not required.")
    document.save(path)


def test_parse_docx_preserves_paragraph_and_table_block_order(tmp_path: Path):
    path = tmp_path / "policy.docx"
    _write_policy(path)

    result = parse(path)

    assert result.error is None
    assert len(result.pages) == 1
    assert result.pages[0].page == 1
    assert result.pages[0].text == result.text
    assert result.pages[0].method == "native"
    assert result.pages[0].quality_score == 1.0
    assert result.text.splitlines() == [
        "# Time off",
        "Type | Minimum",
        "Paid leave | 25 days",
        "Manager approval is not required.",
    ]
    assert result.used_ocr is False
    assert result.quality_score == 1.0


def test_docx_heading_styles_map_to_markdown_headers(tmp_path: Path):
    path = tmp_path / "headings.docx"
    document = Document()
    document.add_heading("Leave policy", level=1)
    document.add_heading("Paid leave", level=2)
    document.add_paragraph("25 days per year.")
    document.add_heading("Sick leave", level=3)
    document.save(path)

    result = parse(path)

    assert result.error is None
    assert result.text.splitlines() == [
        "# Leave policy",
        "## Paid leave",
        "25 days per year.",
        "### Sick leave",
    ]


def test_docx_korean_heading_style_maps_to_markdown_header(tmp_path: Path):
    path = tmp_path / "korean-heading.docx"
    document = Document()
    style = document.styles.add_style("제목 2", WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph("휴가 규정")
    paragraph.style = style
    document.save(path)

    result = parse(path)

    assert result.error is None
    assert result.text == "## 휴가 규정"


def test_docx_heading_level_is_capped_at_six(tmp_path: Path):
    path = tmp_path / "deep-heading.docx"
    document = Document()
    document.add_heading("Deep section", level=7)
    document.save(path)

    result = parse(path)

    assert result.error is None
    assert result.text == "###### Deep section"


def test_corrupt_docx_returns_corrupt_document_error(tmp_path: Path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")

    result = parse(path)

    assert result.error is not None
    assert result.error.code == "corrupt_document"
    assert result.error.message == "DOCX 파일을 읽을 수 없습니다."


def test_empty_docx_returns_native_extraction_empty_error(tmp_path: Path):
    path = tmp_path / "empty.docx"
    Document().save(path)

    result = parse(path)

    assert result.error is not None
    assert result.error.code == "native_extraction_empty"
    assert result.error.message == "DOCX에서 텍스트를 찾지 못했습니다."


def test_docx_with_blank_table_returns_native_extraction_empty_error(tmp_path: Path):
    path = tmp_path / "blank-table.docx"
    document = Document()
    document.add_table(rows=1, cols=2)
    document.save(path)

    result = parse(path)

    assert result.error is not None
    assert result.error.code == "native_extraction_empty"
    assert result.error.message == "DOCX에서 텍스트를 찾지 못했습니다."


def test_docx_merged_table_cell_is_extracted_once(tmp_path: Path):
    path = tmp_path / "merged-cell.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].merge(table.rows[0].cells[1]).text = "Merged value"
    document.save(path)

    result = parse(path)

    assert result.error is None
    assert result.text == "Merged value"


def test_docx_with_malformed_xml_returns_corrupt_document_error(tmp_path: Path):
    path = tmp_path / "malformed.docx"
    Document().save(path)

    with ZipFile(path) as source:
        contents = {
            name: source.read(name)
            for name in source.namelist()
        }
    contents["word/document.xml"] = b"<w:document"
    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, content in contents.items():
            destination.writestr(name, content)

    result = parse(path)

    assert result.error is not None
    assert result.error.code == "corrupt_document"
    assert result.error.message == "DOCX 파일을 읽을 수 없습니다."


def test_docx_extension_is_case_insensitive(tmp_path: Path):
    path = tmp_path / "policy.DOCX"
    _write_policy(path)

    result = parse(path)

    assert result.error is None
    assert result.pages[0].method == "native"
    assert result.text.splitlines() == [
        "# Time off",
        "Type | Minimum",
        "Paid leave | 25 days",
        "Manager approval is not required.",
    ]


def test_docx_parse_accepts_ocr_language_without_changing_extraction(tmp_path: Path):
    path = tmp_path / "policy.docx"
    _write_policy(path)

    result = parse(path, ParseOptions(ocr_language="kor+eng"))

    assert result.error is None
    assert result.text.splitlines() == [
        "# Time off",
        "Type | Minimum",
        "Paid leave | 25 days",
        "Manager approval is not required.",
    ]
